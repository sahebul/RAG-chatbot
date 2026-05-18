import os
import pandas as pd
from docx import Document as DocxDoc
from typing import List
from langchain_core.documents import Document

def load_docx(file_path: str) -> List[Document]:
    doc = DocxDoc(file_path)
    docs = []
    file_name = os.path.basename(file_path)

    full_text = []
    for para in doc.paragraphs:
        if para.text.strip():
            full_text.append(para.text)
    
    if full_text:
        docs.append(Document(
            page_content="\n".join(full_text),
            metadata={
                "source_file": file_name,
                "is_table": False,
                "content_type": "text",
                "page": 0
            }
        ))

    for i, table in enumerate(doc.tables):
        data = []
        for row in table.rows:
            data.append([cell.text.strip() for cell in row.cells])
        if data:
            df = pd.DataFrame(data)
            if not df.empty:
                docs.append(Document(
                    page_content=df.to_markdown(index=False),
                    metadata={
                        "source_file": file_name,
                        "is_table": True,
                        "content_type": "table",
                        "table_index": i,
                        "table_rows": len(df),
                        "table_columns": len(df.columns),
                        "page": 0,
                        "end_page": 0
                    }
                ))
    return docs
